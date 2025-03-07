import os
import pandas as pd
from docx import Document
from datetime import datetime

def extract_data_from_table(docx_path):
    """
    从指定 docx 文件中读取第一个表格，并按照预设的行列坐标提取数据。
    同时扫描文档段落提取测评老师信息（查找“测评老师：”或“测评老师:”）。
    返回一个字典，键为数据名称，值为对应的文本内容。
    """
    try:
        doc = Document(docx_path)
    except Exception as e:
        print(f"打开文件 {docx_path} 出错: {e}")
        return None

    # 提取表格数据
    if not doc.tables:
        print(f"文件 {docx_path} 中未发现表格！")
        return None

    table = doc.tables[0]

    # 定义需要提取的字段与对应的表格坐标（均为 0 索引）
    coords = {
        "儿童姓名": (1, 1),
        "性别": (1, 5),
        "出生日期": (1, 10),
        "测评机构": (3, 1),
        "测评日期": (3, 10),
        "触觉调节-得分": (6, 0),
        "触觉调节_评估等级": (6, 4),
        "前庭觉调节_得分": (6, 7),
        "前庭觉调节_评估等级": (6, 12),
        "本体觉调节_得分": (8, 0),
        "本体觉调节_评估等级": (8, 4),
        "视听觉调节_得分": (8, 7),
        "视听觉调节_评估等级": (8, 12),
        "姿势控制能力_得分": (10, 0),
        "姿势控制能力_评估等级": (10, 4),
        "两侧整合与运用肢体能力_得分": (10, 7),
        "两侧整合与运用肢体能力_评估等级": (10, 12),
        "触觉区辨与精细操作能力_得分": (12, 0),
        "触觉区辨与精细操作能力_评估等级": (12, 4),
        "视听觉区辨能力_得分": (12, 7),
        "视听觉区辨能力_评估等级": (12, 12),
        "感觉敏感_得分": (15, 0),
        "感觉敏感_评估等级": (15, 1),
        "感觉迟钝_得分": (15, 3),
        "感觉迟钝_评估等级": (15, 6),
        "感觉寻求_得分": (15, 9),
        "感觉寻求_评估等级": (15, 11)
    }

    data = {}
    for field, (row, col) in coords.items():
        try:
            data[field] = table.cell(row, col).text.strip()
        except IndexError:
            print(f"文件 {docx_path} 中表格的坐标 ({row}, {col}) 不存在，字段【{field}】未提取！")
            data[field] = ""

    # 提取测评老师信息：扫描文档段落，查找包含“测评老师：”或“测评老师:”的段落
    teacher = ""
    for para in doc.paragraphs:
        if "测评老师" in para.text:
            if "测评老师：" in para.text:
                teacher = para.text.split("测评老师：", 1)[1].strip()
            elif "测评老师:" in para.text:
                teacher = para.text.split("测评老师:", 1)[1].strip()
            if teacher:
                break
    data["测评老师"] = teacher
    return data

def compute_age(birth_date_str, measure_date_str):
    """
    根据出生日期和测评日期（格式均为 'YYYY-MM-DD'）计算满岁年龄。
    若日期格式不正确或任一日期为空，则返回空字符串。
    """
    try:
        birth_date = datetime.strptime(birth_date_str, "%Y-%m-%d")
        measure_date = datetime.strptime(measure_date_str, "%Y-%m-%d")
        age = measure_date.year - birth_date.year - ((measure_date.month, measure_date.day) < (birth_date.month, birth_date.day))
        return age
    except Exception as e:
        print(f"计算年龄时出错：{e}")
        return ""

def process_all_reports(base_folder, output_excel):
    """
    遍历指定的 base_folder（及其所有子文件夹），查找文件名中包含“感统测评报告”的 .docx 文件，
    提取各文件中的数据（包括表格内数据、测评老师信息、测评时年龄和文件名），
    最终将所有数据写入一个 Excel 文件，每个文件对应 Excel 中的一行。
    """
    records = []

    # 遍历大文件夹及所有子文件夹
    for root, dirs, files in os.walk(base_folder):
        for file in files:
            if file.endswith(".docx") and "感统测评报告" in file:
                file_path = os.path.join(root, file)
                print(f"正在处理文件: {file_path}")
                data = extract_data_from_table(file_path)
                if data:
                    # 计算测评时年龄，基于“出生日期”和“测评日期”
                    age = compute_age(data.get("出生日期", ""), data.get("测评日期", ""))
                    data["测评时年龄"] = age
                    # 新增字段：文件名
                    data["文件名"] = file
                    records.append(data)

    if records:
        df = pd.DataFrame(records)
        df.to_excel(output_excel, index=False)
        print(f"所有数据已保存至 {output_excel}")
    else:
        print("未找到符合条件的感统测评报告文件！")

if __name__ == "__main__":
    # 修改此处为大文件夹所在的实际路径
    base_folder = r"/Users/larry/Desktop/长三院儿童康复数据（去重）"
    # 输出 Excel 文件的保存路径
    output_excel = r"/Users/larry/Desktop/长三院儿童康复数据（去重）/集成数据/感统测评.xlsx"
    process_all_reports(base_folder, output_excel)
